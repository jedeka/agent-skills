"""deck.py: scaffold, build, check and export slide decks written as one HTML <section> per slide.

    python3 deck.py new    <deck_dir> --title "T" [--slides main-title,main-content,table] [--footer F]
    python3 deck.py build  <deck_dir>                    # build/<slug>.html, fonts/ and assets/ copied
    python3 deck.py check  <deck_dir>                    # render in Chrome; fit and collision report
    python3 deck.py text   <deck_dir> [--prose]          # plain text of every slide face and its notes
    python3 deck.py png    <deck_dir> [--slides 1,3-5]   # build/png/slide-NN.png, 1920x1080
    python3 deck.py pdf    <deck_dir> [--no-freeze] [--no-pdfcheck]   # build/<slug>.pdf, then pdfcheck
    python3 deck.py pptx   <deck_dir>                    # build/<slug>.pptx (needs python-pptx)
    python3 deck.py inline <deck_dir> [--blobs map.json] # claude.ai Slides files in build/inline/
    python3 deck.py all    <deck_dir>                    # check + pdf (+ pdfcheck) + pptx
    python3 deck.py selftest <out_dir>                   # prove check and pdfcheck catch known failures

Deck layout:
    deck.json         {"title": str, "theme": "technical-swiss", "slides": [ids], "footer": str (optional),
                       "math_fontset": "cm" (optional)}
    slides/<id>.html  exactly one <section>; class "cover" / "section" or data-nofooter = no footer;
                      optional <aside class="notes"> as its last child (speaker notes)
    assets/           images, referenced from slides as assets/<file>

The build inlines templates/<theme>/theme.css, wraps every section in <div class="slide">, inserts
the footer and adds a script that scales the slides to the window width. The footer is the page
number "NN / TOTAL" (counting all slides); deck.json "footer" adds a short label on its left, only
when the user asks for one, and a label that looks like a date is refused. Math written as
<span class="math">TeX</span> or <div class="math-display">TeX</div> is typeset at build time by
mathsvg.py (matplotlib mathtext) into inline SVG; bad TeX stops the build and names the slide, and
TeX left outside a math element (\alpha, s_i, x^{2}) gets a warning.

check renders the deck at the true 1920x1080 canvas and reports, per slide:
    area       a box leaves the content area (section padding); canvas: it leaves the slide
    text       text runs past its own box, or a box's content overflows it (clipped text)
    collision  two objects overlap by more than 2 px (a hairline rule: by more than half its
               thickness): text (measured on its rendered lines), images and SVG (formulas
               included), filled boxes, border rules and pseudo-element shapes such as connector
               arrows; a filled or ruled box never collides with what it contains. Elements inside
               data-layer="overlay" may overlap base content, never each other
    footer     anything other than the footer reaches into the footer band (theme --foot-band)
    image      an image failed to load;  font: text under 24 px (table headers excepted)
    widow      a block of two or more lines ends on a line holding one word (words split at any
               white space; a block with a <br> is skipped). The theme's text-wrap: pretty avoids
               most; fix the rest by cutting or rewording, or bind the last two words with &nbsp;
It also warns, without failing, about every image or other media file in assets/ that no slide face
references (assets/ is copied into build/ whole; references in comments and notes do not count).

text prints each slide's face text and then its speaker notes, under the lines "--- face ---" and
"--- notes ---": the text for rereading the faces as the audience sees them. The face text is what
Chrome lays out, in document order (the templates write it in reading order), one line per block
(heading, paragraph, list item, table row; a <br> starts a new line): capitals where text-transform
sets them, a formula as $TeX$, table cells split by " | ", no footer. Text drawn inside images is not
read. scripts/numtrace.py reads the same text. --prose prints the input for the writing passes
instead: the text as written (no text-transform capitals, so a kicker reads "How to read", not an
acronym "HOW TO READ"), a blank line after every block, so a sentence checker never joins a kicker,
a title and a caption into one sentence, and the headers as Markdown headings ("# slide 01 / 10 ·
id", "## face", "## notes"), which a Markdown-aware checker skips.

pdf first measures the screen layout, then prints with every image and SVG frozen to its measured
size (Chromium's print layout otherwise lets an image that shrank on screen print at its natural
size), then runs pdfcheck.py on the PDF. --no-freeze and --no-pdfcheck switch those steps off.

The PPTX holds one full-bleed picture per slide (13.333 x 7.5 in) plus the speaker notes: it is
faithful to the HTML but not editable. Edit the HTML sources and rebuild.

Chrome is needed for check, text, png, pdf, pptx, inline and selftest: --chrome PATH, else
$CHROME_PATH, else Chrome, Chromium, Edge or chrome-headless-shell on PATH, else a Playwright browser
cache ($PLAYWRIGHT_BROWSERS_PATH or the default). Standard library only, except python-pptx for pptx,
PyMuPDF for the PDF check, matplotlib for math and Pillow for inline's pixel comparison
(scripts/setup.sh installs them all).
"""
import argparse
import glob
import html
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import unquote

import mathsvg      # sibling script; matplotlib itself is imported only when a formula is typeset

SCRIPTS = Path(__file__).resolve().parent
SKILL = SCRIPTS.parent
W, H = 1920, 1080

FIT = """<script>
function fit(){const s=Math.min(1,(window.innerWidth-16)/1920);
document.querySelectorAll('.slide').forEach(d=>{d.style.transform='scale('+s+')';
d.style.marginBottom=(24-1080*(1-s))+'px';});}
window.addEventListener('resize',fit);fit();
</script>"""
# No blank page after the last slide when printing.
EXTRA_CSS = ".slide:last-of-type{break-after:auto;page-break-after:auto}"
# Single-slide render (png): the slide fills the window exactly.
SOLO_CSS = "body{background:#FFFFFF}.slide{margin:0}"
MATH_RE = re.compile(r"""class\s*=\s*["'][^"']*\bmath(-display)?\b""")
ASSET_REF = re.compile(r"""assets/([^"'()\s]+)""")
# files a slide can show; anything else in assets/ (a script, the data behind a chart) is not reported
MEDIA = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".avif", ".svg", ".bmp", ".mp4", ".webm", ".mov",
         ".mp3", ".wav", ".ogg", ".pdf"}
MATH_EL = re.compile(r"""<(span|div)\b[^>]*\bclass\s*=\s*["'][^"']*\bmath(-display)?\b[^>]*>.*?</\1\s*>""", re.S | re.I)
# TeX left as plain text on a slide face: a control word, a braced script, or a one-character subscript
RAW_TEX = re.compile(r"\\[A-Za-z]+|[_^]\{|\b[A-Za-z]+_[A-Za-z0-9]\b")
DATE_RE = re.compile(
    r"\b(19|20)\d\d\b|\b\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}\b|"
    r"\b(jan|feb|mar|apr|jun|jul|aug|sep|sept|oct|nov|dec|january|february|march|april|june|july|"
    r"august|september|october|november|december)\.?\s+\d{1,4}\b|"
    r"\b\d{1,2}\s+(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\b", re.I)

# Injected by check (and by pdf, to measure): writes a JSON report into <pre id="__ppt_report">.
# Per slide: problems, the lowest in-flow box edge, the size of every image/SVG in document order
# (frozen before printing) and the boxes of data-layer="overlay" elements (allowed in pdfcheck).
CHECKER = r"""<script>
(function(){
const W=1920,H=1080,TOL=2;
const BAND=parseFloat(getComputedStyle(document.documentElement).getPropertyValue('--foot-band'))||80;
const GRAPHIC='img,svg,video,canvas,iframe,object,embed';
const BOXY=new Set(['fill','rule','shape']);
function desc(el){
  let s=el.tagName.toLowerCase();
  if(el.classList&&el.classList.length) s+='.'+[...el.classList].join('.');
  const t=(el.textContent||'').replace(/\s+/g,' ').trim();
  if(t&&el.tagName.toLowerCase()!=='svg') s+=' "'+(t.length>48?t.slice(0,45)+'...':t)+'"';
  else if(el.tagName==='IMG') s+=' '+el.getAttribute('src');
  else if(el.getAttribute&&el.getAttribute('aria-label')) s+=' ['+el.getAttribute('aria-label').slice(0,40)+']';
  return s;
}
function shown(el){
  if(el.checkVisibility) return el.checkVisibility({opacityProperty:true,visibilityProperty:true});
  const cs=getComputedStyle(el); return cs.display!=='none'&&cs.visibility!=='hidden'&&parseFloat(cs.opacity)>0;
}
function painted(c){return !!c&&c!=='transparent'&&!/rgba\([^)]*,\s*0\)$/.test(c);}
function graphicsOf(sec){return [...sec.querySelectorAll(GRAPHIC)].filter(g=>
  !(g.parentElement&&g.parentElement.closest('svg'))&&!g.closest('aside.notes,.foot'));}
function owner(el,sec){
  for(let o=el;o&&o!==sec;o=o.parentElement){const d=getComputedStyle(o).display;
    if(d!=='inline'&&d!=='contents') return o;}
  return sec;
}
function textRects(n,rel){
  const t=n.textContent, a=t.search(/\S/), b=t.trimEnd().length;
  const rg=document.createRange(); rg.setStart(n,a); rg.setEnd(n,b);
  const cs=getComputedStyle(n.parentElement), lh=cs.lineHeight==='normal'?0:parseFloat(cs.lineHeight);
  const out=[];
  for(const r of rg.getClientRects()){
    if(r.width<0.5||r.height<0.5) continue;
    // the glyph band: the font's content area, cut to the line box when the line-height is tighter
    let y0=r.top,y1=r.bottom; if(lh&&lh<r.height){const c=(y0+y1)/2; y0=c-lh/2; y1=c+lh/2;}
    out.push(rel({left:r.left,right:r.right,top:y0,bottom:y1}));
  }
  return out;
}
function objects(sec,rel){
  const objs=[], byOwner=new Map();
  const tw=document.createTreeWalker(sec,NodeFilter.SHOW_TEXT);
  while(tw.nextNode()){
    const n=tw.currentNode, p=n.parentElement;
    if(!p||!n.textContent.trim()||p.closest('.foot,aside.notes,script,style,svg')||!shown(p)) continue;
    const o=owner(p,sec); let ob=byOwner.get(o);
    if(!ob){ob={kind:'text',el:o,rects:[],nodes:[]}; byOwner.set(o,ob); objs.push(ob);}
    ob.rects.push(...textRects(n,rel)); ob.nodes.push(n);
  }
  for(const g of graphicsOf(sec)){
    const r=g.getBoundingClientRect();
    if(shown(g)&&r.width>0&&r.height>0) objs.push({kind:'graphic',el:g,rects:[rel(r)]});
  }
  for(const el of sec.querySelectorAll('*')){
    if(el.matches(GRAPHIC)||el.closest('svg,.foot,aside.notes,script,style')||!shown(el)) continue;
    const cs=getComputedStyle(el); if(cs.display==='inline'||cs.display==='contents') continue;
    const r=el.getBoundingClientRect(); if(r.width<=0&&r.height<=0) continue;
    const b=rel(r);
    if(painted(cs.backgroundColor)||cs.backgroundImage!=='none'){objs.push({kind:'fill',el,rects:[b]}); continue;}
    const strips=[];
    for(const [side,cut] of [['Top',w=>({...b,y1:b.y0+w})],['Bottom',w=>({...b,y0:b.y1-w})],
                             ['Left',w=>({...b,x1:b.x0+w})],['Right',w=>({...b,x0:b.x1-w})]]){
      const w=parseFloat(cs['border'+side+'Width']);
      if(w>0&&cs['border'+side+'Style']!=='none'&&painted(cs['border'+side+'Color'])) strips.push(cut(w));
    }
    if(strips.length){objs.push({kind:'rule',el,rects:strips}); continue;}
    const pseudo=['::before','::after'].some(ps=>{const c=getComputedStyle(el,ps).content; return c&&c!=='none'&&c!=='normal';});
    if(pseudo&&!el.children.length&&!el.textContent.trim()) objs.push({kind:'shape',el,rects:[b]});
  }
  return objs;
}
function hit(a,b){
  let best=null;
  for(const r of a.rects) for(const q of b.rects){
    const w=Math.min(r.x1,q.x1)-Math.max(r.x0,q.x0), h=Math.min(r.y1,q.y1)-Math.max(r.y0,q.y0);
    // a hairline rule overlaps by at most its own thickness: scale the tolerance down for it
    const tw=Math.min(TOL,(r.x1-r.x0)/2,(q.x1-q.x0)/2), th=Math.min(TOL,(r.y1-r.y0)/2,(q.y1-q.y0)/2);
    if(w>tw&&h>th&&(!best||w*h>best.w*best.h)) best={w,h,x:Math.max(r.x0,q.x0),y:Math.max(r.y0,q.y0)};
  }
  return best;
}
const overlay=el=>el.closest('[data-layer="overlay"]');
// A block of two or more lines whose last line holds one word, or the tail of a word broken at a
// hyphen or dash. Words split at any white space, &nbsp; included; an inline formula counts as a word.
// A block holding a <br> is skipped: its line breaks are chosen by hand.
function widow(o){
  if(o.el.querySelector('br')) return '';
  const frags=[];
  for(const n of o.nodes){
    const re=/\S+/g; let m;
    while((m=re.exec(n.textContent))){
      const rg=document.createRange(); rg.setStart(n,m.index); rg.setEnd(n,m.index+m[0].length);
      for(const r of rg.getClientRects()) if(r.width>0.5&&r.height>0.5)
        frags.push({t:r.top,b:r.bottom,c:(r.top+r.bottom)/2,w:m[0]});
    }
  }
  for(const g of o.el.querySelectorAll('svg')){
    if(g.parentElement.closest('svg')||owner(g.parentElement,o.el.closest('section'))!==o.el) continue;
    const r=g.getBoundingClientRect(); if(r.width>0.5&&r.height>0.5)
      frags.push({t:r.top,b:r.bottom,c:(r.top+r.bottom)/2,w:g.getAttribute('aria-label')||'formula'});
  }
  if(frags.length<2) return '';
  const last=frags.reduce((a,f)=>f.c>a.c?f:a);
  const tail=frags.filter(f=>f.c>last.t&&f.c<last.b);
  return tail.length===1&&frags.some(f=>f.c<last.t)?tail[0].w:'';
}
function run(){
  const out={slides:[],fonts_failed:[],fonts_loaded:[]};
  document.fonts.forEach(f=>{const n=f.family.replace(/["']/g,'')+' '+f.weight;
    if(f.status==='error') out.fonts_failed.push(n); if(f.status==='loaded') out.fonts_loaded.push(n);});
  document.querySelectorAll('.slide > section').forEach((sec,i)=>{
    const sr=sec.getBoundingClientRect(), cs=getComputedStyle(sec);
    const rel=r=>({x0:r.left-sr.left,y0:r.top-sr.top,x1:r.right-sr.left,y1:r.bottom-sr.top});
    const box={l:parseFloat(cs.paddingLeft),r:sr.width-parseFloat(cs.paddingRight),
               t:parseFloat(cs.paddingTop),b:sr.height-parseFloat(cs.paddingBottom)};
    const rep={index:i+1,content_box:box,lowest:0,objects:0,problems:[],graphics:[],overlays:[]};
    const flagged={area:new Set(),canvas:new Set(),text:new Set()};
    const skip=el=>el.closest('.foot, aside.notes, script, style');
    const pinned=el=>{for(let e=el;e&&e!==sec;e=e.parentElement){
      const p=getComputedStyle(e).position; if(p==='absolute'||p==='fixed') return true;} return false;};
    const under=(el,kind)=>{for(let e=el.parentElement;e&&e!==sec;e=e.parentElement)
      if(flagged[kind].has(e)) return true; return false;};
    const add=(kind,el,msg)=>{flagged[kind]&&flagged[kind].add(el);
      if(!flagged[kind]||!under(el,kind)) rep.problems.push({kind:kind,el:desc(el),msg:msg});};
    sec.querySelectorAll('*').forEach(el=>{
      if(skip(el)||el.closest('svg')&&el.tagName.toLowerCase()!=='svg') return;
      const r=el.getBoundingClientRect();
      const x0=r.left-sr.left, x1=r.right-sr.left, y0=r.top-sr.top, y1=r.bottom-sr.top;
      const visible=r.width>0&&r.height>0&&getComputedStyle(el).visibility!=='hidden';
      if(visible){
        if(x0<-1||y0<-1||x1>W+1||y1>H+1)
          add('canvas',el,`box ${Math.round(x0)},${Math.round(y0)} to ${Math.round(x1)},${Math.round(y1)} leaves the ${W}x${H} canvas`);
        else if(!pinned(el)){
          rep.lowest=Math.max(rep.lowest,Math.round(y1));
          if(x0<box.l-1||x1>box.r+1||y1>box.b+1)
            add('area',el,`box ${Math.round(x0)},${Math.round(y0)} to ${Math.round(x1)},${Math.round(y1)} leaves the content area ${box.l}-${box.r} x ${box.t}-${box.b}`);
        }
        // vertical slack: tight line-heights let glyphs poke out of their own line box
        const slack=Math.max(2,0.3*parseFloat(getComputedStyle(el).fontSize));
        if(el.clientWidth>0&&el.clientHeight>0&&!el.closest('svg')&&
           (el.scrollWidth>el.clientWidth+2||el.scrollHeight>el.clientHeight+slack))
          add('text',el,`content ${el.scrollWidth}x${el.scrollHeight} overflows its ${el.clientWidth}x${el.clientHeight} box`);
      }
      if(el.tagName==='IMG'&&(!el.complete||el.naturalWidth===0))
        add('image',el,'image failed to load: '+el.getAttribute('src'));
      const own=[...el.childNodes].some(n=>n.nodeType===3&&n.textContent.trim())||el.matches('.math,.math-display');
      if(own&&visible&&!el.closest('th')){
        let fs=parseFloat(getComputedStyle(el).fontSize);
        if(el.closest('svg')&&el.getScreenCTM){const m=el.getScreenCTM(); if(m) fs*=Math.hypot(m.a,m.b);}
        if(fs<23.5) add('font',el,`font-size ${fs.toFixed(1)} px < 24 px`);
      }
    });
    const objs=objects(sec,rel); rep.objects=objs.length;
    // R7: text past its own box, every pair of objects, the footer band
    for(const o of objs){
      if(o.kind!=='text'||o.el===sec) continue;
      const b=rel(o.el.getBoundingClientRect());
      const over=Math.max(0,...o.rects.map(r=>Math.max(r.x1-b.x1,b.x0-r.x0)));
      if(over>TOL&&!flagged.text.has(o.el)) add('text',o.el,`text runs ${Math.round(over)} px past its ${Math.round(b.x1-b.x0)} px wide box`);
    }
    for(const o of objs){
      if(o.kind!=='text'||o.el===sec) continue;
      const w=widow(o); if(w) add('widow',o.el,`the last line holds one word: "${w.length>30?w.slice(0,27)+'...':w}"`);
    }
    for(let a=0;a<objs.length;a++) for(let c=a+1;c<objs.length;c++){
      const p=objs[a], q=objs[c];
      if(p.el===q.el||(BOXY.has(p.kind)&&p.el.contains(q.el))||(BOXY.has(q.kind)&&q.el.contains(p.el))) continue;
      if(!!overlay(p.el)!==!!overlay(q.el)) continue;
      const h=hit(p,q);
      if(h) rep.problems.push({kind:'collision',el:desc(p.el)+' ('+p.kind+')',
        msg:`overlaps ${desc(q.el)} (${q.kind}) by ${Math.round(h.w)} x ${Math.round(h.h)} px at (${Math.round(h.x)}, ${Math.round(h.y)})`});
    }
    for(const o of objs){
      const y=Math.max(...o.rects.map(r=>r.y1));
      if(y>H-BAND+TOL) rep.problems.push({kind:'footer',el:desc(o.el)+' ('+o.kind+')',
        msg:`reaches y ${Math.round(y)}; the footer band starts at y ${H-BAND}`});
    }
    for(const g of graphicsOf(sec)){const r=g.getBoundingClientRect();
      rep.graphics.push(r.width>0&&r.height>0?[+r.width.toFixed(2),+r.height.toFixed(2)]:null);}
    for(const o of sec.querySelectorAll('[data-layer="overlay"]')){const b=rel(o.getBoundingClientRect());
      rep.overlays.push([b.x0-TOL,b.y0-TOL,b.x1+TOL,b.y1+TOL].map(v=>Math.round(v)));}
    out.slides.push(rep);
  });
  const pre=document.createElement('pre'); pre.id='__ppt_report';
  pre.textContent=JSON.stringify(out); document.body.appendChild(pre);
}
window.addEventListener('load',()=>document.fonts.ready.then(()=>setTimeout(run,50)));
})();
</script>"""

# Injected by pdf: sets every image/SVG to the size the screen layout gave it, before printing.
FREEZE = """<script>
(function(){const F=__SIZES__;
const G='img,svg,video,canvas,iframe,object,embed';
document.querySelectorAll('.slide > section').forEach((sec,i)=>{
  const gs=[...sec.querySelectorAll(G)].filter(g=>!(g.parentElement&&g.parentElement.closest('svg'))&&!g.closest('aside.notes,.foot'));
  (F[i]||[]).forEach((s,k)=>{const g=gs[k]; if(!g||!s) return;
    for(const [p,v] of [['box-sizing','border-box'],['width',s[0]+'px'],['height',s[1]+'px'],['min-width','0'],
        ['min-height','0'],['max-width','none'],['max-height','none'],['flex','none']]) g.style.setProperty(p,v,'important');});
});})();
</script>"""

# Injected by text: the face text of every slide as laid out (innerText), formulas as $TeX$, without
# the footer and the notes; with __PROSE__ true, text-transform is dropped first (the case as written).
TEXTER = r"""<script>
(function(){
function run(){
  const out=[];
  if(__PROSE__){const st=document.createElement('style');
    st.textContent='.slide > section, .slide > section *{text-transform:none !important}';
    document.head.appendChild(st);}
  document.querySelectorAll('.slide > section').forEach(sec=>{
    sec.querySelectorAll('.foot, aside.notes, script, style').forEach(e=>e.remove());
    sec.querySelectorAll('.math svg[aria-label], .math-display svg[aria-label]').forEach(g=>{
      if(!g.parentElement.closest('svg')) g.replaceWith(document.createTextNode('$'+g.getAttribute('aria-label')+'$'));});
    out.push(sec.innerText);
  });
  const pre=document.createElement('pre'); pre.id='__deck_text';
  pre.textContent=JSON.stringify(out); document.body.appendChild(pre);
}
window.addEventListener('load',()=>document.fonts.ready.then(()=>setTimeout(run,50)));
})();
</script>"""

CHROME_NAMES = ["google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "chrome",
                "chrome-headless-shell", "msedge"]
CHROME_GLOBS = ["chromium_headless_shell-*/chrome-*/chrome-headless-shell", "chromium-*/chrome-*/chrome",
                "chromium-*/chrome-mac/Chromium.app/Contents/MacOS/Chromium"]
MAC_APPS = ["/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Chromium.app/Contents/MacOS/Chromium"]


class DeckError(Exception):
    """A user-facing error: printed without a traceback."""


# ---------------------------------------------------------------- deck sources

def load_deck(deck_dir: Path) -> dict:
    path = deck_dir / "deck.json"
    if not path.is_file():
        raise DeckError(f"{path} not found (scaffold one with: deck.py new {deck_dir} --title ...)")
    deck = json.loads(path.read_text(encoding="utf-8"))
    for key in ("title", "slides"):
        if key not in deck:
            raise DeckError(f"{path}: missing key '{key}'")
    if not deck["slides"]:
        raise DeckError(f"{path}: 'slides' is empty")
    deck.setdefault("theme", "technical-swiss")
    deck.setdefault("footer", "")
    deck.setdefault("math_fontset", "cm")
    if deck["math_fontset"] not in mathsvg.FONTSETS:
        raise DeckError(f'{path}: unknown math_fontset "{deck["math_fontset"]}"; '
                        f"choose one of: {mathsvg.FONTSET_HELP}")
    if DATE_RE.search(deck["footer"]):
        raise DeckError(f'{path}: footer "{deck["footer"]}" looks like it holds a date. The footer is the '
                        "page number plus, only when asked, a short label; never a date")
    return deck


def theme_dir(theme: str) -> Path:
    d = SKILL / "templates" / theme
    if not (d / "theme.css").is_file():
        known = sorted(p.parent.name for p in (SKILL / "templates").glob("*/theme.css"))
        raise DeckError(f"theme '{theme}' not found; available: {', '.join(known) or 'none'}")
    return d


def foot_band(deck: dict) -> float:
    """Height of the footer band in px, from the theme's --foot-band (default 80)."""
    m = re.search(r"--foot-band\s*:\s*([\d.]+)px", (theme_dir(deck["theme"]) / "theme.css").read_text())
    return float(m[1]) if m else 80.0


def slug(title: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-") or "deck"


def read_slides(deck_dir: Path, deck: dict) -> list:
    """[(id, html, notes_text)], validating one <section> per file and typesetting math."""
    out = []
    for sid in deck["slides"]:
        path = deck_dir / "slides" / f"{sid}.html"
        if not path.is_file():
            raise DeckError(f"slide '{sid}' listed in deck.json but {path} is missing")
        src = path.read_text(encoding="utf-8")
        n = len(re.findall(r"<section\b", src, re.I))
        if n != 1:
            raise DeckError(f"{path}: expected exactly one <section>, found {n}")
        for tag in ("<html", "<head", "<style"):
            if tag in src.lower():
                print(f"warning: {path.name} contains {tag}>; slides hold only a <section>",
                      file=sys.stderr)
        for ref in ASSET_REF.findall(src):
            if not (deck_dir / "assets" / ref).exists():
                print(f"warning: {path.name} references assets/{ref}, which does not exist",
                      file=sys.stderr)
        notes = notes_text(src)
        face = MATH_EL.sub(" ", re.sub(r"<!--.*?-->|<aside\b.*?</aside>", " ", src, flags=re.S | re.I))
        raw = sorted(set(RAW_TEX.findall(html.unescape(re.sub(r"<[^>]*>", " ", face)))))
        if raw:
            print(f"warning: {path.name} shows TeX as plain text ({', '.join(raw)}); write it as "
                  '<span class="math">TeX</span>', file=sys.stderr)
        if MATH_RE.search(re.sub(r"<!--.*?-->", "", src, flags=re.S)):   # math markup in a comment is no math
            src = typeset(src, sid, deck["math_fontset"])
        out.append((sid, src.strip(), notes))
    return out


def typeset(src: str, sid: str, fontset: str) -> str:
    if importlib.util.find_spec("matplotlib") is None:
        raise DeckError(f"slide {sid} has math markup; typesetting it needs matplotlib "
                        "(pip install matplotlib, or run scripts/setup.sh)")
    try:
        return mathsvg.replace_math(src, fontset)
    except mathsvg.MathError as e:
        raise DeckError(f"slide {sid}: {e}")


def notes_text(src: str) -> str:
    m = re.search(r"""<aside\b[^>]*class=["'][^"']*\bnotes\b[^"']*["'][^>]*>(.*?)</aside>""", src,
                  re.S | re.I)
    if not m:
        return ""
    txt = re.sub(r"<br\s*/?>|</p>|</li>|</div>", "\n", m[1], flags=re.I)
    txt = html.unescape(re.sub(r"<[^>]+>", "", txt))
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in txt.splitlines()]
    return "\n".join(ln for ln in lines if ln)


def with_footer(src: str, footer: str, page: int, total: int) -> str:
    """Insert the footer before </section> unless the section is a cover/section/nofooter slide."""
    attrs = re.search(r"<section\b([^>]*)>", src, re.I)[1]
    cls = re.search(r"""class\s*=\s*["']([^"']*)["']""", attrs)
    classes = cls[1].split() if cls else []
    if "cover" in classes or "section" in classes or re.search(r"\bdata-nofooter\b", attrs):
        return src
    label = f'<span class="foot-label">{html.escape(footer)}</span>' if footer else ""
    foot = f'<div class="foot">{label}<span class="foot-num">{page:02d} / {total:02d}</span></div>\n'
    i = src.lower().rfind("</section>")
    return src[:i] + foot + src[i:]


def page_html(deck: dict, slides: list, fit: bool, extra_css: str = "", tail: str = "",
              base: str = "", only: int = 0) -> str:
    """The deck as one HTML document. only=N keeps slide N (1-based) alone, footer numbering intact."""
    css = (theme_dir(deck["theme"]) / "theme.css").read_text(encoding="utf-8")
    total = len(slides)
    body = []
    for i, (_, src, _) in enumerate(slides, 1):
        if only and i != only:
            continue
        body.append(f'<div class="slide">\n{with_footer(src, deck["footer"], i, total)}\n</div>')
    base_tag = f'<base href="{base}">\n' if base else ""
    return ('<!doctype html>\n<html lang="en"><head><meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
            f'{base_tag}<title>{html.escape(deck["title"])}</title>\n'
            f'<style>\n{css}\n{EXTRA_CSS}\n{extra_css}</style></head>\n<body>\n'
            + "\n".join(body) + "\n" + (FIT if fit else "") + tail + "\n</body></html>\n")


# ---------------------------------------------------------------- commands

def cmd_new(a) -> int:
    deck_dir = Path(a.deck_dir)
    if (deck_dir / "deck.json").exists() and not a.force:
        raise DeckError(f"{deck_dir}/deck.json exists; pass --force to overwrite")
    if DATE_RE.search(a.footer):
        raise DeckError(f'--footer "{a.footer}" looks like a date; footers never carry dates')
    tdir = theme_dir(a.theme)
    templates = {re.sub(r"^\d+-", "", p.stem): p for p in sorted((tdir / "slides").glob("*.html"))}
    if not templates:
        raise DeckError(f"no slide templates in {tdir / 'slides'}")
    types = [t.strip() for t in a.slides.split(",") if t.strip()] if a.slides else list(templates)
    unknown = [t for t in types if re.sub(r"^\d+-", "", t) not in templates]
    if unknown:
        raise DeckError(f"unknown slide type(s): {', '.join(unknown)}\n"
                        f"available: {', '.join(templates)}")
    (deck_dir / "slides").mkdir(parents=True, exist_ok=True)
    (deck_dir / "assets").mkdir(exist_ok=True)
    ids = []
    for n, t in enumerate(types, 1):
        t = re.sub(r"^\d+-", "", t)
        sid = f"{n:02d}-{t}"
        src = templates[t].read_text(encoding="utf-8")
        (deck_dir / "slides" / f"{sid}.html").write_text(src, encoding="utf-8")
        ids.append(sid)
        for ref in ASSET_REF.findall(src):
            hits = [p for p in (tdir / "assets" / ref, tdir / "slides" / "assets" / ref) if p.is_file()]
            if hits:
                (deck_dir / "assets" / ref).parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(hits[0], deck_dir / "assets" / ref)
            else:
                print(f"warning: template {t} references assets/{ref}, not found in {tdir}",
                      file=sys.stderr)
    deck = {"title": a.title, "theme": a.theme, "slides": ids}
    if a.footer:
        deck["footer"] = a.footer
    (deck_dir / "deck.json").write_text(json.dumps(deck, indent=2, ensure_ascii=False) + "\n",
                                        encoding="utf-8")
    print(f"scaffolded {deck_dir}: {len(ids)} slides ({', '.join(ids)})")
    return 0


def cmd_build(a, quiet: bool = False) -> tuple:
    """Write build/<slug>.html; return (deck, slides) for the commands that render next."""
    deck_dir = Path(a.deck_dir)
    deck = load_deck(deck_dir)
    slides = read_slides(deck_dir, deck)
    build = deck_dir / "build"
    build.mkdir(exist_ok=True)
    for name, src in (("fonts", theme_dir(deck["theme"]) / "fonts"), ("assets", deck_dir / "assets")):
        if src.is_dir():
            shutil.copytree(src, build / name, dirs_exist_ok=True)
    out = build / f"{slug(deck['title'])}.html"
    out.write_text(page_html(deck, slides, fit=True), encoding="utf-8")
    if not quiet:
        print(f"wrote {out} ({len(slides)} slides)")
    return deck, slides


def render_json(a, deck: dict, slides: list, script: str, pre_id: str, what: str):
    """Render the deck at 1920x1080 with `script` injected; return the JSON it writes into <pre id=pre_id>."""
    probe = Path(a.deck_dir) / "build" / f"_{what}.html"
    probe.write_text(page_html(deck, slides, fit=False, tail=script), encoding="utf-8")
    try:
        dom = chrome(a, ["--dump-dom", probe.resolve().as_uri()], capture=True)
    finally:
        probe.unlink(missing_ok=True)
    m = re.search(f'<pre id="{pre_id}">(.*?)</pre>', dom, re.S)
    if not m:
        raise DeckError(f"the {what} script produced no output (Chrome did not finish loading the page; "
                        "try a larger --budget)")
    return json.loads(html.unescape(m[1]))


def render_report(a, deck: dict, slides: list) -> dict:
    """Render the deck at 1920x1080 with the checker injected; return its JSON report."""
    return render_json(a, deck, slides, CHECKER, "__ppt_report", "check")


def unused_assets(deck_dir: Path, deck: dict) -> list:
    """Media files in assets/ that no slide face references (comments and notes left out)."""
    root = deck_dir / "assets"
    if not root.is_dir():
        return []
    used = set()
    for sid in deck["slides"]:
        src = (deck_dir / "slides" / f"{sid}.html").read_text(encoding="utf-8")
        face = re.sub(r"<!--.*?-->|<aside\b.*?</aside>", " ", src, flags=re.S | re.I)
        used |= {unquote(html.unescape(r)) for r in ASSET_REF.findall(face)}
    files = {f.relative_to(root).as_posix() for f in root.rglob("*")
             if f.is_file() and f.suffix.lower() in MEDIA and not f.name.startswith(".")}
    return sorted(files - used)


def slide_texts(a) -> list:
    """Per slide {"n", "id", "face", "notes"}: the face text as Chrome lays it out (see the module
    docstring, text) and the speaker notes. Builds the deck first."""
    deck, slides = cmd_build(a, quiet=True)
    script = TEXTER.replace("__PROSE__", "true" if getattr(a, "prose", False) else "false")
    faces = render_json(a, deck, slides, script, "__deck_text", "text")
    out = []
    for n, ((sid, _, notes), face) in enumerate(zip(slides, faces), 1):
        face = re.sub(r"[\u00a0\u2007\u2009\u202f]", " ", face).replace("\t", " | ")
        lines = [re.sub(r" {2,}", " ", ln).strip() for ln in face.splitlines()]
        out.append({"n": n, "id": sid, "face": "\n".join(ln for ln in lines if ln), "notes": notes})
    return out


def cmd_text(a) -> int:
    texts = slide_texts(a)
    for t in texts:
        if a.prose:                     # Markdown: checkers skip headings; a blank line ends each block
            print(f"# slide {t['n']:02d} / {len(texts):02d} · {t['id']}\n\n## face\n")
            print("\n\n".join(t["face"].splitlines()) or "(no text)")
            print("\n## notes\n")
            print(t["notes"] or "(no notes)")
        else:
            print(f"=== slide {t['n']:02d} / {len(texts):02d} · {t['id']} ===")
            print("--- face ---")
            print(t["face"] or "(no text)")
            print("--- notes ---")
            print(t["notes"] or "(no notes)")
        print()
    return 0


def cmd_check(a) -> tuple:
    """Print the fit and collision report; return (exit code, report)."""
    deck, slides = cmd_build(a)
    rep = render_report(a, deck, slides)
    bad = 0
    print(f"\n{'#':>3}  {'slide':<32} {'lowest y':>8}  {'objects':>7}  problems")
    for (sid, _, _), s in zip(slides, rep["slides"]):
        limit = s["content_box"]["b"]
        flag = "" if s["lowest"] <= limit else f"  (limit {limit:g})"
        print(f"{s['index']:>3}  {sid:<32} {s['lowest']:>8}  {s['objects']:>7}  {len(s['problems']) or '-'}{flag}")
    for (sid, _, _), s in zip(slides, rep["slides"]):
        for p in s["problems"]:
            bad += 1
            print(f"  slide {s['index']:02d} [{p['kind']}] {p['el']}: {p['msg']}")
    for f in rep["fonts_failed"]:
        bad += 1
        print(f"  font failed to load: {f}")
    for f in unused_assets(Path(a.deck_dir), deck):
        print(f"  warning: assets/{f} is not referenced by any slide (build/ still copies it)")
    print(f"\nfonts loaded: {', '.join(sorted(rep['fonts_loaded'])) or 'none (system fallback)'}")
    hit = sum(1 for s in rep["slides"] if s["problems"])
    print(f"\n{bad} problem(s); {hit} of {len(slides)} slide(s) affected" if bad
          else f"\nall {len(slides)} slides fit; no collisions")
    return (1 if bad else 0), rep


def parse_range(spec: str, total: int) -> list:
    if not spec:
        return list(range(1, total + 1))
    nums = []
    for part in spec.split(","):
        lo, _, hi = part.strip().partition("-")
        nums += range(int(lo), int(hi or lo) + 1)
    wrong = [n for n in nums if not 1 <= n <= total]
    if wrong:
        raise DeckError(f"slide number(s) {wrong} outside 1-{total}")
    return sorted(set(nums))


def cmd_png(a, quiet: bool = False) -> tuple:
    """Render the chosen slides; return (paths of the PNGs, slides)."""
    deck_dir = Path(a.deck_dir)
    deck, slides = cmd_build(a, quiet=quiet)
    nums = parse_range(getattr(a, "slides", ""), len(slides))
    out_dir = deck_dir / "build" / "png"
    out_dir.mkdir(exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="_render-", dir=deck_dir / "build"))

    def shot(n: int) -> Path:
        page = tmp / f"slide-{n:02d}.html"
        page.write_text(page_html(deck, slides, fit=False, extra_css=SOLO_CSS, base="../", only=n),
                        encoding="utf-8")
        png = out_dir / f"slide-{n:02d}.png"
        png.unlink(missing_ok=True)
        chrome(a, [f"--screenshot={png.resolve()}", page.resolve().as_uri()],
               profile=tmp / f"profile-{n}")
        if not png.is_file():
            raise DeckError(f"Chrome wrote no screenshot for slide {n}")
        size = png_size(png)
        if size != (W, H):
            print(f"warning: {png.name} is {size[0]}x{size[1]}, expected {W}x{H}", file=sys.stderr)
        return png

    try:
        with ThreadPoolExecutor(max_workers=min(4, os.cpu_count() or 1)) as pool:
            pngs = list(pool.map(shot, nums))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if not quiet:
        print(f"wrote {len(pngs)} PNG(s) to {out_dir}")
    return pngs, slides


def cmd_pdf(a, rep: dict = None) -> int:
    """Print the deck (images frozen at their screen size), then check the PDF with pdfcheck.py."""
    deck_dir = Path(a.deck_dir)
    deck, slides = cmd_build(a, quiet=True)
    check_pdf = not getattr(a, "no_pdfcheck", False)
    freeze = not getattr(a, "no_freeze", False)
    if check_pdf and importlib.util.find_spec("pymupdf") is None:
        raise DeckError("the PDF check needs PyMuPDF (pip install pymupdf, or run scripts/setup.sh); "
                        "pass --no-pdfcheck to print without checking")
    if rep is None and (freeze or check_pdf):
        rep = render_report(a, deck, slides)
    tail = ""
    if freeze:
        tail = FREEZE.replace("__SIZES__", json.dumps([s["graphics"] for s in rep["slides"]]))
    build = deck_dir / "build"
    page = build / "_print.html"
    page.write_text(page_html(deck, slides, fit=False, tail=tail), encoding="utf-8")
    out = build / f"{slug(deck['title'])}.pdf"
    out.unlink(missing_ok=True)
    try:
        chrome(a, [f"--print-to-pdf={out.resolve()}", "--no-pdf-header-footer",
                   page.resolve().as_uri()])
    finally:
        page.unlink(missing_ok=True)
    if not out.is_file():
        raise DeckError("Chrome wrote no PDF")
    pages = len(re.findall(rb"/Type\s*/Page\b", out.read_bytes()))
    note = "" if pages == len(slides) else f" (expected {len(slides)}: check for overflowing slides)"
    frozen = sum(1 for s in rep["slides"] for g in s["graphics"] if g) if freeze else 0
    how = f"{frozen} image(s)/SVG(s) frozen at their screen size" if freeze else "images not frozen"
    print(f"wrote {out} ({pages} pages{note}; {how})")
    if not check_pdf:
        print("PDF not checked (--no-pdfcheck)")
        return 0
    cmd = [sys.executable, str(SCRIPTS / "pdfcheck.py"), str(out), "--pages", str(len(slides)),
           "--band", f"{foot_band(deck):g}", "--label", deck["footer"]]
    allow = {str(s["index"]): s["overlays"] for s in rep["slides"] if s["overlays"]}
    allow_file = build / "_overlays.json"
    if allow:
        allow_file.write_text(json.dumps({"pages": allow}), encoding="utf-8")
        cmd += ["--allow", str(allow_file)]
    try:
        r = subprocess.run(cmd, capture_output=True, text=True)
    finally:
        allow_file.unlink(missing_ok=True)
    print(r.stdout.rstrip())
    if r.returncode not in (0, 1):
        raise DeckError(f"pdfcheck failed: {r.stderr.strip()[-600:]}")
    return r.returncode


def cmd_pptx(a) -> int:
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        raise DeckError("pptx export needs python-pptx: pip install python-pptx (or run scripts/setup.sh)")
    deck_dir = Path(a.deck_dir)
    deck = load_deck(deck_dir)
    a.slides = ""
    pngs, slides = cmd_png(a, quiet=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 13.333 x 7.5 in
    blank = prs.slide_layouts[6]
    for png, (_, _, notes) in zip(pngs, slides):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    cp = prs.core_properties
    cp.title, cp.author, cp.last_modified_by, cp.comments = deck["title"], "", "", ""
    out = deck_dir / "build" / f"{slug(deck['title'])}.pptx"
    prs.save(out)
    print(f"wrote {out} ({len(pngs)} slides, pictures with speaker notes)")
    return 0


def cmd_inline(a) -> int:
    import inline
    return inline.run(a)


def cmd_all(a) -> int:
    rc, rep = cmd_check(a)
    print()
    rc_pdf = cmd_pdf(a, rep)
    cmd_pptx(a)
    if rc or rc_pdf:
        print("\n" + ("check" if rc else "") + (" and " if rc and rc_pdf else "")
              + ("the PDF check" if rc_pdf else "") + " reported problems: fix them and rebuild")
    return 1 if rc or rc_pdf else 0


# ---------------------------------------------------------------- selftest

FIXTURE_SLIDE = """<section>
  <div class="head"><p class="kicker">Fixture · {label}</p><h2>{title}</h2></div>
  {body}
  <aside class="notes">Fixture for deck.py selftest; the content is illustrative, not data.</aside>
</section>"""
NUMBER = '<p style="font-size:72px;font-weight:700;white-space:nowrap" class="red">35.3%</p>'
EXPLAIN = ('<p class="s">Fixture text beside a 72 px bold number. The number only reproduces the width '
           'of the figure in the incident; it is not data.</p>')
IMG = '<img src="assets/photo.png" alt="4:3 placeholder" style="width:860px;height:500px">'
PIN = 'style="position:absolute;left:40px;top:40px"'
# deck name -> slides as (label, title, body, the problem kind check must report, or None for a clean slide)
FIXTURES = {
    "r7-incident": [  # 2026-10-19: a 72 px bold number in a column sized by a character estimate
        ("beside", "A large number beside its explanation",
         f'<div class="panel" style="display:grid;grid-template-columns:180px 1fr;gap:24px;align-items:center">'
         f'{NUMBER}{EXPLAIN}</div>', "collision")],
    "r7-fixed": [  # the two fixes R7 names
        ("explicit width", "A large number beside its explanation",
         f'<div class="panel" style="display:grid;grid-template-columns:300px 1fr;gap:24px;align-items:center">'
         f'{NUMBER}{EXPLAIN}</div>', None),
        ("number above", "A large number above its explanation",
         f'<div class="panel stack c8" style="gap:12px">{NUMBER}{EXPLAIN}</div>', None)],
    "r7-kinds": [  # one object kind per slide, then a one-word last line
        ("text on image", "A label set on an image",
         f'<div class="c6" style="position:relative">{IMG}<p class="s" {PIN}>A label on the image</p></div>', "collision"),
        ("overlay", "A label marked as an overlay",
         f'<div class="c6" style="position:relative">{IMG}<p class="s" data-layer="overlay" {PIN}>A label on the '
         f'image</p></div>', None),
        ("overlay clash", "Two overlay labels on each other",
         f'<div class="c6" style="position:relative">{IMG}<p class="s" data-layer="overlay" {PIN}>First label on '
         f'top</p><p class="s" data-layer="overlay" style="position:absolute;left:60px;top:50px">Second label</p></div>',
         "collision"),
        ("rule", "A hairline rule through text",
         '<div style="position:relative"><p>A paragraph that a rule crosses through its first line.</p>'
         '<div style="position:absolute;left:0;right:0;top:20px;border-top:2px solid #111"></div></div>', "collision"),
        ("connector", "A connector under text",
         '<div class="flow"><div class="node"><h3>Stage one</h3></div><div class="arrow"></div>'
         '<div class="node"><h3 style="margin-left:-60px">Stage two</h3></div></div>', "collision"),
        ("footer band", "A note pinned into the footer band",
         '<p class="meta" style="position:absolute;left:88px;bottom:30px">Source line pinned too low</p>', "footer"),
        ("table cell", "A table cell wider than its column",
         '<table style="table-layout:fixed;width:900px"><tr><th style="width:200px">Name</th><th>Value</th></tr>'
         '<tr><td style="white-space:nowrap">a-very-long-identifier-name</td><td>0.95</td></tr></table>', "collision"),
        ("formula", "A formula over text",
         '<p>Body text above the formula, a full line long enough to meet it.</p>'
         '<div class="math-display" style="margin-top:-60px">\\sum_{i=1}^{n} x_i = 1</div>', "collision"),
        ("panel", "A filled panel over text",
         '<p>Body text that a panel slides up over from below.</p>'
         '<div class="panel" style="margin-top:-70px;width:600px"><p class="s">Panel</p></div>', "collision"),
        ("clean", "Image, text and formula side by side",
         f'<div class="row"><div class="c6">{IMG}</div><div class="c6 stack"><p>Text beside the image.</p>'
         f'<div class="math-display">\\alpha = 0.05</div></div></div>', None),
        ("widow", "A paragraph whose last line is one word",
         '<p class="c4" style="text-wrap:wrap">Plain wrapping, turned back on for this fixture, leaves one word '
         'alone.</p>', "widow")],
}
PRINT_SLIDE = """<section>
  <div class="head"><p class="kicker">Print fixture</p><h2>A figure that shrinks to fit on screen</h2></div>
  <div style="display:grid;grid-template-columns:5fr 7fr;grid-template-rows:minmax(0,1fr);gap:24px;flex:1;min-height:0">
    <p>The image is a 4:3 placeholder in a column that is shorter than its natural height.</p>
    <div class="fig" style="min-height:0">
      <img src="assets/photo.png" alt="4:3 placeholder" style="flex:1 1 auto;min-height:0;width:100%">
      <p class="cap">Caption under the figure</p>
    </div>
  </div>
  <div class="panel"><p>A filled bar under the figure row</p></div>
</section>"""


def write_fixture(root: Path, name: str, slides: list) -> None:
    d = root / name
    shutil.rmtree(d, ignore_errors=True)
    (d / "slides").mkdir(parents=True)
    (d / "assets").mkdir()
    shutil.copy2(SKILL / "templates" / "technical-swiss" / "assets" / "placeholder-photo-a.png",
                 d / "assets" / "photo.png")
    ids = []
    for i, (sid, src) in enumerate(slides, 1):
        ids.append(f"{i:02d}-{slug(sid)}")
        (d / "slides" / f"{ids[-1]}.html").write_text(src, encoding="utf-8")
    (d / "deck.json").write_text(json.dumps({"title": name, "slides": ids}, indent=2))


def cmd_selftest(a) -> int:
    """Build fixture decks that must fail and pass, run check and pdf on them, compare verdicts."""
    root = Path(a.out_dir)
    if root.exists() and any(root.iterdir()) and not (root / ".deck-selftest").exists():
        raise DeckError(f"{root} is not empty; give an empty or new folder")
    if importlib.util.find_spec("matplotlib") is None or importlib.util.find_spec("pymupdf") is None:
        raise DeckError("selftest needs matplotlib and PyMuPDF (run scripts/setup.sh)")
    root.mkdir(parents=True, exist_ok=True)
    (root / ".deck-selftest").write_text("written by deck.py selftest\n")
    for name, slides in FIXTURES.items():
        write_fixture(root, name, [(lab, FIXTURE_SLIDE.format(label=lab, title=t, body=b)) for lab, t, b, _ in slides])
    write_fixture(root, "print-shrink", [("shrink", PRINT_SLIDE)])

    opts = (["--chrome", a.chrome] if a.chrome else []) + ["--budget", str(a.budget)]
    cases = [  # (args, expected exit code, tokens the output must hold ("!" = must not), per-slide kinds, what it shows)
        (["check", "r7-incident"], 1, ["[text]"], FIXTURES["r7-incident"],
         "the incident: the number overlaps its text and runs past its column"),
        (["check", "r7-fixed"], 0, ["no collisions", "warning: assets/photo.png is not referenced"],
         FIXTURES["r7-fixed"], "both fixes R7 names pass: explicit width with slack, number above text "
         "(the unused fixture image is named in a warning, which does not fail the check)"),
        (["check", "r7-kinds"], 1, ["!is not referenced"], FIXTURES["r7-kinds"],
         "every object kind and a one-word last line are caught; an overlay label and a clean layout pass"),
        (["check", "print-shrink"], 0, ["no collisions"], None,
         "the print fixture's screen layout is clean, so only the PDF check can catch it"),
        (["pdf", "print-shrink", "--no-freeze"], 1, ["PDF FAIL", "[footer]"], None,
         "printed without freezing, the figure takes its natural size and pdfcheck flags the spill"),
        (["pdf", "print-shrink"], 0, ["PDF OK"], None,
         "printed with the figure frozen at its screen size, the PDF passes"),
        (["pdf", "r7-kinds"], 1, ["page 1: [image]", "page 3: [text]", "page 6: [footer]", "page 9: [panel]",
                                  "!page 2:", "!page 10:"], None,
         "pdfcheck finds the text, image, panel and footer cases in print; the overlay label is allowed"),
    ]
    ok = 0
    for args, want, tokens, kinds, what in cases:
        cmd = [sys.executable, str(Path(__file__).resolve()), args[0], str(root / args[1])] + args[2:] + opts
        r = subprocess.run(cmd, capture_output=True, text=True)
        found = {}
        for n, kind in re.findall(r"^\s+slide (\d+) \[(\w+)\]", r.stdout, re.M):
            found.setdefault(int(n), set()).add(kind)
        wrong = [f"slide {i:02d} expected {k or 'no problem'}, got {', '.join(sorted(found.get(i, []))) or 'none'}"
                 for i, (_, _, _, k) in enumerate(kinds or [], 1)
                 if (k is None and found.get(i)) or (k is not None and k not in found.get(i, set()))]
        good = r.returncode == want and not wrong and all(
            (t[1:] not in r.stdout) if t.startswith("!") else (t in r.stdout) for t in tokens)
        ok += good
        print(f"{'PASS' if good else 'FAIL'}  deck.py {' '.join(args)}  (exit {r.returncode}, expected {want}): {what}")
        for line in wrong + [ln.strip() for ln in (r.stdout + r.stderr).splitlines()
                             if re.search(r"^\s+slide \d+ \[|^page \d+: \[|PDF (OK|FAIL)|slides fit|problem\(s\)|^error", ln)]:
            print(f"      | {line}")
    print(f"\nselftest {'passed' if ok == len(cases) else 'FAILED'}: {ok} of {len(cases)} as expected "
          f"(fixtures in {root})")
    return 0 if ok == len(cases) else 1


# ---------------------------------------------------------------- Chrome

def find_chrome(explicit: str = "") -> str:
    for c in (explicit, os.environ.get("CHROME_PATH", "")):
        if c:
            if shutil.which(c) or Path(c).is_file():
                return shutil.which(c) or c
            raise DeckError(f"Chrome not found at {c}")
    for name in CHROME_NAMES:
        if shutil.which(name):
            return shutil.which(name)
    for app in MAC_APPS:
        if Path(app).is_file():
            return app
    roots = [Path.home() / ".cache" / "ms-playwright", Path.home() / "Library" / "Caches" / "ms-playwright"]
    if os.environ.get("PLAYWRIGHT_BROWSERS_PATH"):
        roots.insert(0, Path(os.environ["PLAYWRIGHT_BROWSERS_PATH"]))
    for root in roots:
        for pattern in CHROME_GLOBS:
            hits = sorted(glob.glob(str(root / pattern)))
            if hits:
                return hits[-1]
    raise DeckError("no Chrome or Chromium found. Provide one with --chrome PATH or CHROME_PATH=..., "
                    "install google-chrome or chromium on PATH, or run scripts/setup.sh "
                    "(installs a headless Chromium through Playwright).")


def chrome(a, args: list, capture: bool = False, profile: Path = None) -> str:
    exe = find_chrome(getattr(a, "chrome", "") or "")
    headless = "--headless" if "headless-shell" in Path(exe).name else "--headless=new"
    cmd = [exe, headless, "--no-sandbox", "--disable-gpu", "--hide-scrollbars",
           "--force-device-scale-factor=1", f"--window-size={W},{H}",
           f"--virtual-time-budget={getattr(a, 'budget', 8000)}",
           "--allow-file-access-from-files", "--no-first-run", "--no-default-browser-check",
           "--disable-extensions", "--mute-audio"]
    if profile:
        cmd.append(f"--user-data-dir={profile}")
    try:
        r = subprocess.run(cmd + args, capture_output=True, text=True, timeout=180)
    except subprocess.TimeoutExpired:
        raise DeckError(f"Chrome timed out: {' '.join(args)}")
    if r.returncode != 0:
        raise DeckError(f"Chrome exited {r.returncode}: {r.stderr.strip()[-800:]}")
    return r.stdout if capture else ""


def png_size(path: Path) -> tuple:
    head = path.read_bytes()[:24]
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")


# ---------------------------------------------------------------- CLI

def main(argv=None) -> int:
    sys.modules.setdefault("deck", sys.modules[__name__])   # inline.py imports this module as `deck`
    sys.stdout.reconfigure(line_buffering=True)      # keep stdout and stderr lines in order
    chromeopts = argparse.ArgumentParser(add_help=False)
    chromeopts.add_argument("--chrome", default="", help="path to Chrome/Chromium (else CHROME_PATH, PATH)")
    chromeopts.add_argument("--budget", type=int, default=8000,
                            help="Chrome virtual-time budget in ms (time for fonts and images to load)")
    common = argparse.ArgumentParser(add_help=False, parents=[chromeopts])
    common.add_argument("deck_dir", help="deck folder holding deck.json, slides/ and assets/")
    printing = argparse.ArgumentParser(add_help=False)
    printing.add_argument("--no-freeze", action="store_true",
                          help="print without freezing images at their screen size")
    printing.add_argument("--no-pdfcheck", action="store_true", help="do not run pdfcheck.py on the PDF")
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = p.add_subparsers(dest="cmd", required=True)
    n = sub.add_parser("new", parents=[common], help="scaffold a deck from the theme's templates")
    n.add_argument("--title", required=True)
    n.add_argument("--slides", default="", help="comma-separated slide types (default: all templates)")
    n.add_argument("--footer", default="",
                   help="optional short label left of the page number, only when asked; never a date")
    n.add_argument("--theme", default="technical-swiss")
    n.add_argument("--force", action="store_true", help="overwrite an existing deck.json")
    sub.add_parser("build", parents=[common], help="write build/<slug>.html")
    sub.add_parser("check", parents=[common], help="build, render, report fit, overlap and small type")
    x = sub.add_parser("text", parents=[common], help="print each slide's face text and speaker notes")
    x.add_argument("--prose", action="store_true",
                   help="the text as written, a blank line after each block, Markdown headers: for the writing passes")
    g = sub.add_parser("png", parents=[common], help="render slides to build/png/slide-NN.png")
    g.add_argument("--slides", default="", help="1-based slide numbers, e.g. 1,3-5 (default: all)")
    sub.add_parser("pdf", parents=[common, printing], help="write build/<slug>.pdf and check it")
    sub.add_parser("pptx", parents=[common], help="write build/<slug>.pptx (pictures + notes)")
    i = sub.add_parser("inline", parents=[common],
                       help="export to the claude.ai Slides format in build/inline/ (see inline.py)")
    i.add_argument("--out", default="", help="output folder (default <deck_dir>/build/inline)")
    i.add_argument("--blobs", default="", help='JSON {"assets/x.png": "/_blob/<id>", ...} of uploaded images')
    i.add_argument("--no-verify", action="store_true", help="skip the pixel comparison")
    i.add_argument("--threshold", type=float, default=1.0,
                   help="largest mean absolute pixel difference allowed per slide (0-255, default 1.0)")
    sub.add_parser("all", parents=[common, printing], help="check + pdf (+ pdfcheck) + pptx")
    t = sub.add_parser("selftest", parents=[chromeopts],
                       help="build fixture decks and confirm check and pdfcheck catch known failures")
    t.add_argument("out_dir", help="an empty folder for the fixture decks")
    a = p.parse_args(argv)
    try:
        if a.cmd == "build":
            cmd_build(a)
            return 0
        if a.cmd == "png":
            cmd_png(a)
            return 0
        if a.cmd == "check":
            return cmd_check(a)[0]
        return {"new": cmd_new, "text": cmd_text, "pdf": cmd_pdf, "pptx": cmd_pptx, "inline": cmd_inline, "all": cmd_all,
                "selftest": cmd_selftest}[a.cmd](a)
    except DeckError as e:
        print(f"error: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
